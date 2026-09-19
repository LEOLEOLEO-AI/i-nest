"""Build the final manuscript and its evidence/engineering appendices."""

from pathlib import Path
import re
import subprocess

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


SOURCE = Path(__file__).resolve().parent
OUTPUT = Path("D:/Obsidian/Output/18_iNEST_Final_Theory_20260917")
PANDOC = Path("C:/Users/LEO/pandoc/pandoc-3.6.3/pandoc.exe")
PARTS = ["01_Final_Theory.md", "06_SEI_Engineering_Protocol.md",
         "05_Numerical_Checks.md", "02_Claim_Passport.md", "03_Evidence_Integration.md"]


def set_font(style, size, bold=False):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    reference = Document()
    section = reference.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = reference.styles["Normal"]
    set_font(normal, 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in [("Title", 23), ("Subtitle", 12), ("Heading 1", 17),
                       ("Heading 2", 13), ("Heading 3", 11.5)]:
        set_font(reference.styles[name], size, name.startswith("Heading"))
        reference.styles[name].font.color.rgb = RGBColor.from_string("203D43")
    for name in ["TOC 1", "TOC 2"]:
        style = reference.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_font(style, 9.5)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0
    reference_path = OUTPUT / "03_Word_Reference.docx"
    reference.save(reference_path)
    target = OUTPUT / "01_Final_Theory.docx"
    # Pandoc's Word math conversion drops LaTeX tags; render numbers as math text.
    render_text = "\n\n".join((SOURCE / name).read_text(encoding="utf-8") for name in PARTS)
    render_text = re.sub(r"\\tag\{(\d+)\}", lambda m: r"\qquad\text{(" + m[1] + ")}", render_text)
    render_source = OUTPUT / "07_Render_Source.md"
    render_source.write_text(render_text, encoding="utf-8")
    subprocess.run([str(PANDOC), str(render_source),
                    "--standalone", "--toc", "--toc-depth=1", "--reference-doc", str(reference_path),
                    "-o", str(target)], check=True)
    doc = Document(target)
    for node in doc.element.xpath(".//w:t"):
        if node.text == "Table of Contents":
            node.text = "目录"
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 1":
            paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.widow_control = True
    for table in doc.tables:
        table.autofit = True
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in table.rows:
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        footer.add_run("iNEST-C | 2026-09-18 | ").font.size = Pt(8)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    doc.save(target)
    print(target)


if __name__ == "__main__":
    main()
