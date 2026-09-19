"""Build a styled Word edition with verification and evidence appendices."""

from pathlib import Path
import subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path(__file__).resolve().parent
OUTPUT = Path("D:/Obsidian/Output/17_iNEST_Reconstruction_20260917")
PANDOC = Path("C:/Users/LEO/pandoc/pandoc-3.6.3/pandoc.exe")


def font(style, size, bold=False):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    reference = Document()
    section = reference.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    font(reference.styles["Normal"], 10.5)
    reference.styles["Normal"].paragraph_format.space_after = Pt(6)
    reference.styles["Normal"].paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 23), ("Subtitle", 12), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        font(reference.styles[name], size, name.startswith("Heading"))
        reference.styles[name].font.color.rgb = RGBColor.from_string("203D43")
    reference_path = OUTPUT / "03_Word_Reference.docx"
    reference.save(reference_path)
    target = OUTPUT / "01_Theory_Audit_and_Reconstruction.docx"
    files = [SOURCE / name for name in ("01_Theory_Audit_and_Reconstruction.md",
             "02_Verification_Plan.md", "03_Evidence_Ledger.md", "08_Numerical_Checks.md")]
    subprocess.run([str(PANDOC), *map(str, files), "--standalone", "--toc", "--toc-depth=2",
                    "--reference-doc", str(reference_path), "-o", str(target)], check=True)
    document = Document(target)
    for node in document.element.xpath(".//w:t"):
        if node.text == "Table of Contents":
            node.text = "目录"
    for paragraph in document.paragraphs:
        if paragraph.text == "Table of Contents":
            paragraph.text = "目录"
        if paragraph.style.name == "Heading 1":
            paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.widow_control = True
    for table in document.tables:
        table.autofit = True
        header = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        header.append(repeat)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        run = footer.add_run("iNEST-R | 2026-09-17 | ")
        run.font.size = Pt(8)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    document.save(target)
    print(target)


if __name__ == "__main__":
    main()
